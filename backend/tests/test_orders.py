"""Orders and invoices: the lines and their totals (worked out on the server, per VAT rate), numbering, the status, and the invoice
file made from the built-in layout, from a Word template and from a PDF used as letter paper."""
import io
from datetime import date, timedelta
from decimal import Decimal

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from app import invoice_docs
from tests.test_auth_api import H, post
from tests.test_domain_api import make, owner, sheet_payload      # noqa: F401  (fixture)


def line(description="Vaas", quantity=1, unit_price="12.10", vat_pct="21", **extra):
    return {"description": description, "quantity": quantity, "unit_price": unit_price, "vat_pct": vat_pct, **extra}


def new_order(owner, lines=None, **extra):
    body = {"customer_name": "Jan Jansen", "customer_number": "K9999", "subject": "Rack-houders voor oktober", "address": "Kerkstraat 1", "postal_code": "1234 AB", "city": "Utrecht", "country": "Nederland",
            "email": "jan@example.nl", "order_date": "2026-10-01", "delivery_date": "2026-10-15",
            "lines": [line()] if lines is None else lines, **extra}
    res = owner.post("/api/orders", json=body, headers=H)
    assert res.status_code == 201, res.get_json()
    return res.get_json()


def business(owner, **extra):
    body = {"name": "Flux76 Prints", "address": "Industrieweg 5", "postal_code": "3500 AA", "city": "Utrecht", "email": "info@flux76.nl",
            "kvk": "12345678", "vat_id": "NL001234567B01", "iban": "NL91ABNA0417164300", "bic": "ABNANL2A", "website": "flux76.nl",
            "phone": "06 12345678", "note": "Bedankt voor je bestelling!", "payment_days": 14, "invoice_language": "nl", **extra}
    res = owner.put("/api/preferences", json={"business": body}, headers=H)
    assert res.status_code == 200, res.get_json()
    return res.get_json()["business"]


def pdf_text(content: bytes) -> str:
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)


def upload(owner, name, content, **form):
    return owner.post("/api/invoice-template", data={"file": (io.BytesIO(content), name), **form}, headers=H,
                      content_type="multipart/form-data")


def docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts += [cell.text for row in table.rows for cell in row.cells]
    return "\n".join(parts)


# --- calculation numbers --------------------------------------------------------------------------------------

def test_calculations_have_a_number_to_find_them_by(owner):
    first = make(owner, "/api/projects", sheet_payload())
    second = make(owner, "/api/projects", {**sheet_payload(), "name": "Tweede"})
    assert (first["number"], second["number"]) == (1, 2)
    assert {p["number"] for p in owner.get("/api/projects").get_json()} == {1, 2}
    owner.delete(f"/api/projects/{first['id']}", headers=H)
    assert make(owner, "/api/projects", sheet_payload())["number"] == 3          # a number is never given out twice


# --- totals ---------------------------------------------------------------------------------------------------

def test_the_vat_is_taken_out_of_the_prices_per_rate(owner):
    order = new_order(owner, [line("A", 2, "12.10", "21"), line("B", 1, "10.90", "9"), line("C", 1, "5.00", "0")])
    totals = order["totals"]
    assert totals["line_totals"] == ["24.20", "10.90", "5.00"]
    assert totals["vat"] == [{"pct": "0", "net": "5.00", "vat": "0.00", "gross": "5.00"},
                             {"pct": "9", "net": "10.00", "vat": "0.90", "gross": "10.90"},
                             {"pct": "21", "net": "20.00", "vat": "4.20", "gross": "24.20"}]
    assert (totals["total_excl"], totals["total_vat"], totals["total"]) == ("35.00", "5.10", "40.10")


def test_lines_with_the_same_rate_are_added_up_before_the_vat_is_taken_out(owner):
    order = new_order(owner, [line("A", 1, "0.99", "21"), line("B", 1, "0.99", "21"), line("C", 1, "0.99", "21")])
    (row,) = order["totals"]["vat"]
    assert row["gross"] == "2.97" and row["vat"] == "0.52" and row["net"] == "2.45"             # not 3 x 0.17 + ...
    assert Decimal(row["net"]) + Decimal(row["vat"]) == Decimal(row["gross"])


def test_the_totals_endpoint_gives_what_saving_will_store(owner):
    lines = [line("A", 3, "4.99", "21"), line("B", 1, "20.00", "9")]
    shown = post(owner, "/api/orders/totals", {"lines": lines}).get_json()
    assert shown == new_order(owner, lines)["totals"]
    assert post(owner, "/api/orders/totals", {"lines": [line(quantity=0)]}).status_code == 400


# --- validation -----------------------------------------------------------------------------------------------

@pytest.mark.parametrize("change", [
    {"customer_name": ""}, {"status": "klaar"}, {"order_date": "morgen"}, {"delivery_date": "31-12-2026"}, {"email": "geen-adres"},
    {"lines": "veel"}, {"lines": [line(quantity=0)]}, {"lines": [line(quantity=1.5)]}, {"lines": [line(quantity="x")]},
    {"lines": [line(unit_price="-1")]}, {"lines": [line(unit_price="abc")]}, {"lines": [line(vat_pct="101")]},
    {"lines": [line(description="")]}, {"lines": [line(project_id="niet-een-uuid")]}, {"lines": [line()] * 201},
    {"customer_name": "x" * 121},
])
def test_an_order_that_makes_no_sense_is_refused(owner, change):
    body = {"customer_name": "Jan", "order_date": "2026-10-01", "lines": [line()], **change}
    assert owner.post("/api/orders", json=body, headers=H).status_code == 400


def test_an_order_may_start_without_lines_and_with_only_a_name(owner):
    order = owner.post("/api/orders", json={"customer_name": "Jan"}, headers=H).get_json()
    assert order["lines"] == [] and order["totals"]["total"] == "0.00" and order["status"] == "nieuw"
    assert order["order_date"] == date.today().isoformat() and order["number"] == 1


def test_every_order_endpoint_needs_a_session(client):
    for method, path in (("get", "/api/orders"), ("post", "/api/orders"), ("get", "/api/orders/" + "0" * 8 + "-0000-0000-0000-" + "0" * 12),
                         ("post", "/api/orders/totals"), ("get", "/api/invoice-template"), ("get", "/api/invoice-template/sample")):
        assert getattr(client, method)(path, headers=H).status_code == 401, path


# --- lines linked to calculations -------------------------------------------------------------------------------

def test_a_line_remembers_its_calculation_and_survives_its_removal(owner):
    project = make(owner, "/api/projects", sheet_payload())
    order = new_order(owner, [line("Rack-houder", 2, "4.50", "21", project_id=project["id"])])
    assert order["lines"][0]["project_id"] == project["id"] and order["lines"][0]["project_number"] == project["number"]
    owner.delete(f"/api/projects/{project['id']}", headers=H)
    kept = owner.get(f"/api/orders/{order['id']}").get_json()
    assert kept["lines"][0]["project_id"] is None and kept["lines"][0]["project_number"] == project["number"]
    assert kept["lines"][0]["description"] == "Rack-houder" and kept["totals"]["total"] == "9.00"


def test_a_line_keeps_the_price_it_was_made_with(owner):
    project = make(owner, "/api/projects", sheet_payload())
    order = new_order(owner, [line("X", 1, "9.99", "21", project_id=project["id"])])
    owner.put(f"/api/projects/{project['id']}", json={**sheet_payload(), "margin_pct": "300"}, headers=H)
    assert owner.get(f"/api/orders/{order['id']}").get_json()["lines"][0]["unit_price"] == "9.99"


# --- list, search, status, editing --------------------------------------------------------------------------------

def test_the_list_filters_by_status_and_searches_by_name_number_and_city(owner):
    a = new_order(owner, customer_name="Anna de Vries", city="Zwolle")
    b = new_order(owner, customer_name="Bram Bakker", city="Utrecht", order_date="2026-11-01")
    assert [o["number"] for o in owner.get("/api/orders").get_json()] == [b["number"], a["number"]]         # newest first
    assert [o["customer_name"] for o in owner.get("/api/orders?q=vries").get_json()] == ["Anna de Vries"]
    assert [o["customer_name"] for o in owner.get("/api/orders?q=utrecht").get_json()] == ["Bram Bakker"]
    assert [o["customer_name"] for o in owner.get(f"/api/orders?q=%23{a['number']}").get_json()] == ["Anna de Vries"]
    assert owner.patch(f"/api/orders/{a['id']}/status", json={"status": "verstuurd"}, headers=H).status_code == 200
    assert [o["customer_name"] for o in owner.get("/api/orders?status=verstuurd").get_json()] == ["Anna de Vries"]
    assert owner.get("/api/orders?status=onzin").status_code == 400
    assert owner.patch(f"/api/orders/{a['id']}/status", json={"status": "onzin"}, headers=H).status_code == 400
    listed = owner.get("/api/orders").get_json()[0]
    assert listed["total"] == "12.10" and listed["line_count"] == 1


def test_an_order_can_be_edited_and_deleted(owner):
    order = new_order(owner)
    body = {**order, "customer_name": "Piet", "lines": [line("Nieuw", 4, "2.50", "0")]}
    saved = owner.put(f"/api/orders/{order['id']}", json=body, headers=H).get_json()
    assert saved["customer_name"] == "Piet" and [l["description"] for l in saved["lines"]] == ["Nieuw"] and saved["totals"]["total"] == "10.00"
    assert owner.delete(f"/api/orders/{order['id']}", headers=H).status_code == 200
    assert owner.get(f"/api/orders/{order['id']}").status_code == 404
    assert owner.put(f"/api/orders/{order['id']}", json=body, headers=H).status_code == 404


def test_the_status_can_be_set_to_every_stage(owner):
    order = new_order(owner)
    for status in ("verwerkt", "verstuurd", "afgehandeld", "nieuw"):
        assert owner.patch(f"/api/orders/{order['id']}/status", json={"status": status}, headers=H).get_json()["status"] == status
        assert owner.get(f"/api/orders/{order['id']}").get_json()["status"] == status


# --- the company details ---------------------------------------------------------------------------------------

def test_company_details_are_validated_and_given_defaults(owner):
    saved = owner.put("/api/preferences", json={"business": {"name": "Flux"}}, headers=H).get_json()["business"]
    assert saved["name"] == "Flux" and saved["payment_days"] == 14 and saved["invoice_language"] == "nl" and saved["iban"] == ""
    for bad in ("x", None, {"payment_days": -1}, {"payment_days": "veel"}, {"invoice_language": "fr"}, {"name": "x" * 121}):
        assert owner.put("/api/preferences", json={"business": bad}, headers=H).status_code == 400, bad


# --- the invoice: numbering ---------------------------------------------------------------------------------------

def test_an_order_without_lines_cannot_be_invoiced(owner):
    order = new_order(owner, [])
    assert post(owner, f"/api/orders/{order['id']}/invoice").status_code == 409


def test_an_invoice_number_is_given_once_and_counted_per_year(owner, app):
    first, second = new_order(owner), new_order(owner)
    year = date.today().year
    one = post(owner, f"/api/orders/{first['id']}/invoice")
    assert one.status_code == 200 and one.mimetype == "application/pdf"
    assert f'filename="Factuur-{year}-0001.pdf"' in one.headers["Content-Disposition"]
    again = post(owner, f"/api/orders/{first['id']}/invoice")
    assert f"{year}-0001" in again.headers["Content-Disposition"]                                  # the same number again
    assert f"{year}-0002" in post(owner, f"/api/orders/{second['id']}/invoice").headers["Content-Disposition"]
    saved = owner.get(f"/api/orders/{first['id']}").get_json()
    assert saved["invoice_number"] == f"{year}-0001" and saved["invoice_date"] == date.today().isoformat()
    with app.extensions["pool"].connection() as conn:                                              # a new year starts again at 1
        conn.execute("UPDATE orders SET invoice_number = NULL, invoice_date = %s WHERE id = %s", (date(year - 1, 12, 30), second["id"]))
    assert f"{year - 1}-0001" in post(owner, f"/api/orders/{second['id']}/invoice").headers["Content-Disposition"]


def test_the_invoice_date_can_be_changed_once_there_is_an_invoice(owner):
    order = new_order(owner)
    owner.put(f"/api/orders/{order['id']}", json={**order, "invoice_date": "2026-01-01"}, headers=H)
    assert owner.get(f"/api/orders/{order['id']}").get_json()["invoice_date"] is None            # no invoice yet: nothing to date
    post(owner, f"/api/orders/{order['id']}/invoice")
    saved = owner.get(f"/api/orders/{order['id']}").get_json()
    changed = owner.put(f"/api/orders/{order['id']}", json={**saved, "invoice_date": "2026-01-01"}, headers=H).get_json()
    assert changed["invoice_date"] == "2026-01-01" and changed["invoice_number"] == saved["invoice_number"]


# --- the invoice: built-in layout -------------------------------------------------------------------------------

def test_the_built_in_invoice_follows_the_example_layout(owner):
    business(owner)
    order = new_order(owner, [line("Rack-houder", 2, "12.10", "21"), line("Verzending", 1, "5.00", "0")])
    text = pdf_text(post(owner, f"/api/orders/{order['id']}/invoice").data)
    year = date.today().year
    due = (date.today() + timedelta(days=14)).strftime("%d/%m/%Y")
    for expected in ("Flux76 Prints", "Industrieweg 5", "3500 AA Utrecht", "Btw-nummer: NL001234567B01", "KVK-nummer: 12345678", "Tel.: 06 12345678",
                     "E-mail: info@flux76.nl", "Website: flux76.nl", "IBAN: NL91ABNA0417164300", "BIC: ABNANL2A",
                     "Jan Jansen", "Kerkstraat 1", "1234 AB Utrecht",
                     f"Factuurdatum: {date.today().strftime('%d/%m/%Y')}", f"Vervaldatum: {due}", "Klantnummer: K9999",
                     f"Factuur: {year}-0001", "Betreft: Rack-houders voor oktober",
                     "Aantal", "Beschrijving", "Bedrag", "excl. btw", "incl. btw", "Rack-houder", "Verzending", "€ 20,00", "€ 24,20", "€ 5,00", "21% Btw", "0% Btw",
                     "Totaalbedrag excl. btw", "€ 25,00", "Btw hoog (21%)", "€ 4,20", "Totaalbedrag incl. btw", "€ 29,20",
                     f"Gelieve dit bedrag van € 29,20 te betalen voor: {due} onder vermelding van {year}-0001.", "Bedankt voor je bestelling!"):
        assert expected in text, expected


def test_a_subject_that_was_not_filled_in_says_which_order_it_is(owner):
    business(owner)
    order = new_order(owner, subject="")
    text = pdf_text(post(owner, f"/api/orders/{order['id']}/invoice").data)
    assert f"Betreft: Bestelling {order['number']}" in text


def test_the_invoice_can_be_in_english(owner):
    business(owner, invoice_language="en")
    order = new_order(owner, subject="")
    text = pdf_text(post(owner, f"/api/orders/{order['id']}/invoice").data)
    for expected in ("Invoice: ", "Subject: Order 1", "Invoice date: ", "Due date: ", "Customer number: K9999", "Description", "Total incl. VAT",
                     "VAT standard rate (21%)", "Please pay this amount of € 12.10 by", "VAT no.: NL001234567B01", "CoC no.: 12345678"):
        assert expected in text, expected
    assert "Factuur" not in text and "Beschrijving" not in text


def test_an_order_without_vat_shows_only_the_total(owner):
    order = new_order(owner, [line("Vrijgesteld", 3, "5.00", "0")])
    text = pdf_text(post(owner, f"/api/orders/{order['id']}/invoice").data)
    assert "€ 15,00" in text and "Totaalbedrag excl. btw" not in text and "Btw hoog" not in text and "Totaalbedrag" in text


def test_a_long_order_continues_on_a_second_page_with_the_header_repeated(owner):
    order = new_order(owner, [line(f"Model nummer {i}", 1, "1.00", "21") for i in range(80)])
    reader = PdfReader(io.BytesIO(post(owner, f"/api/orders/{order['id']}/invoice").data))
    assert len(reader.pages) >= 2 and "Beschrijving" in reader.pages[1].extract_text() and "Model nummer 79" in reader.pages[-1].extract_text()


def test_markup_in_a_name_is_not_taken_for_markup(owner):
    order = new_order(owner, [line("<b>Vet</b> & <i>schuin</i>")], customer_name="A & B <script>")
    text = pdf_text(post(owner, f"/api/orders/{order['id']}/invoice").data)
    assert "<b>Vet</b> & <i>schuin</i>" in text and "A & B <script>" in text


# --- the invoice: Word template ---------------------------------------------------------------------------------

def lines_table(docx_bytes):
    """The table that holds the lines (its first cell says Aantal), among the layout tables of the template."""
    return next(t for t in Document(io.BytesIO(docx_bytes)).tables if t.rows[0].cells[0].text.strip() == "Aantal")


def test_the_sample_template_is_a_word_file_in_the_example_layout(owner):
    res = owner.get("/api/invoice-template/sample")
    assert res.status_code == 200 and "wordprocessingml" in res.mimetype
    text = docx_text(res.data)
    for name in ("logo", "betaal_qr", "factuurnummer", "factuurdatum", "vervaldatum", "betreft", "klant.naam", "klant.adres", "klant.postcode_plaats",
                 "klant.nummer", "regel.aantal", "regel.omschrijving", "regel.bedrag_excl", "regel.bedrag_incl", "regel.btw", "totaal_excl",
                 "btw_labels", "btw_bedragen", "totaal", "betaaltekst", "bedrijf.naam", "bedrijf.btw_nummer", "bedrijf.kvk", "bedrijf.telefoon",
                 "bedrijf.email", "bedrijf.website", "bedrijf.iban", "bedrijf.bic", "bedrijf.tekst"):
        assert "{{%s}}" % name in text, name
    assert [c.text.strip() for c in lines_table(res.data).rows[0].cells] == ["Aantal", "Beschrijving", "Bedrag\nexcl. btw", "Bedrag\nincl. btw"]
    assert owner.get("/api/invoice-template").get_json()["placeholders"] == list(invoice_docs.PLACEHOLDERS)


def test_every_placeholder_that_is_offered_is_actually_filled_in(owner):
    business(owner)
    plain = [n for n in invoice_docs.PLACEHOLDERS if not n.startswith("regel.")]
    template = make_docx(paragraphs=["{{%s}}" % n for n in plain], rows=[["{{%s}}" % n for n in invoice_docs.PLACEHOLDERS if n.startswith("regel.")]])
    assert upload(owner, "alles.docx", template).status_code == 201
    order = new_order(owner, [line("Rack", 2, "12.10", "21")], notes="Graag in doos")
    text = docx_text(post(owner, f"/api/orders/{order['id']}/invoice").data)
    assert "{{" not in text and "Graag in doos" in text and "NL001234567B01" in text and "Nederland" in text and "jan@example.nl" in text


def test_an_order_is_filled_into_the_sample_template_with_one_row_per_line(owner):
    business(owner)
    sample = owner.get("/api/invoice-template/sample").data
    res = upload(owner, "mijn factuur.docx", sample)
    assert res.status_code == 201 and res.get_json()["template"]["kind"] == "docx" and res.get_json()["template"]["filename"] == "mijn factuur.docx"
    order = new_order(owner, [line("Rack-houder", 2, "12.10", "21"), line("Verzending", 1, "5.00", "0")])
    made = post(owner, f"/api/orders/{order['id']}/invoice")
    assert made.status_code == 200 and "wordprocessingml" in made.mimetype
    year = date.today().year
    assert f'Factuur-{year}-0001.docx' in made.headers["Content-Disposition"]
    text = docx_text(made.data)
    assert "{{" not in text
    for expected in ("Flux76 Prints", "Industrieweg 5", "KVK-nummer: 12345678", "IBAN: NL91ABNA0417164300", "BIC: ABNANL2A", "Website: flux76.nl",
                     f"Factuur: {year}-0001", "Betreft: Rack-houders voor oktober", "Klantnummer: K9999", "Jan Jansen", "Kerkstraat 1", "1234 AB Utrecht",
                     "Rack-houder", "Verzending", "€ 20,00", "€ 24,20", "€ 5,00", "21% Btw", "Totaalbedrag incl. btw", "€ 29,20", "€ 25,00",
                     "Btw hoog (21%)", "€ 4,20", "Gelieve dit bedrag van € 29,20"):
        assert expected in text, expected
    table = lines_table(made.data)
    assert [row.cells[0].text for row in table.rows] == ["Aantal", "2", "1"]                       # header + one row per line
    assert [row.cells[1].text for row in table.rows][1:] == ["Rack-houder", "Verzending"]
    assert len(Document(io.BytesIO(made.data)).inline_shapes) == 1                                      # the payment QR; there is no logo yet


def make_docx(paragraphs=(), rows=(), runs=None):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if runs:
        paragraph = doc.add_paragraph()
        for piece in runs:
            paragraph.add_run(piece)
    if rows:
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for row, values in zip(table.rows, rows):
            for cell, value in zip(row.cells, values):
                cell.text = value
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_a_placeholder_that_word_split_over_several_runs_is_still_filled_in(owner):
    template = make_docx(runs=["Aan: {{kl", "ant.na", "am}} te ", "{{klant.plaats}}"], rows=[["{{regel.omschrijving}}", "{{regel.totaal}}"]])
    assert upload(owner, "t.docx", template).status_code == 201
    order = new_order(owner)
    text = docx_text(post(owner, f"/api/orders/{order['id']}/invoice").data)
    assert "Aan: Jan Jansen te Utrecht" in text


def test_a_placeholder_that_does_not_exist_is_left_visible_so_a_typo_shows(owner):
    template = make_docx(paragraphs=["{{klant.naam}} / {{klant.nmaa}}"], rows=[["{{regel.omschrijving}}"]])
    assert upload(owner, "t.docx", template).status_code == 201
    text = docx_text(post(owner, f"/api/orders/{new_order(owner)['id']}/invoice").data)
    assert "Jan Jansen / {{klant.nmaa}}" in text


def test_a_word_template_needs_a_row_for_the_lines(owner):
    res = upload(owner, "t.docx", make_docx(paragraphs=["{{klant.naam}}"]))
    assert res.status_code == 400 and "{{regel.omschrijving}}" in res.get_json()["error"]
    assert owner.get("/api/invoice-template").get_json()["template"] is None


def test_a_header_and_footer_are_filled_in_too(owner):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Kop {{bedrijf.naam}}"
    doc.sections[0].footer.paragraphs[0].text = "Voet {{factuurnummer}}"
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "{{regel.omschrijving}}"
    out = io.BytesIO()
    doc.save(out)
    business(owner)
    assert upload(owner, "t.docx", out.getvalue()).status_code == 201
    made = post(owner, f"/api/orders/{new_order(owner)['id']}/invoice").data
    doc = Document(io.BytesIO(made))
    assert doc.sections[0].header.paragraphs[0].text == "Kop Flux76 Prints"
    assert doc.sections[0].footer.paragraphs[0].text == f"Voet {date.today().year}-0001"


# --- the invoice: PDF as letter paper -------------------------------------------------------------------------------

def letter_paper(text="BRIEFPAPIER Flux76", size=(595.27, 841.89)):
    out = io.BytesIO()
    pdf = canvas.Canvas(out, pagesize=size)
    pdf.drawString(60, size[1] - 40, text)
    pdf.save()
    return out.getvalue()


def test_the_invoice_is_drawn_on_the_letter_paper_and_the_paper_stays(owner):
    business(owner)
    res = upload(owner, "briefpapier.pdf", letter_paper(), margin_top_mm="60", margin_bottom_mm="30")
    assert res.status_code == 201 and res.get_json()["template"]["margin_top_mm"] == 60
    order = new_order(owner)
    made = post(owner, f"/api/orders/{order['id']}/invoice")
    assert made.mimetype == "application/pdf"
    reader = PdfReader(io.BytesIO(made.data))
    text = reader.pages[0].extract_text()
    assert "BRIEFPAPIER Flux76" in text and "Jan Jansen" in text and "Btw-nummer" not in text and "€ 12,10" in text
    assert "Industrieweg" not in text                       # the paper carries the company details, so they are not printed twice
    assert float(reader.pages[0].mediabox.width) == pytest.approx(595.27, abs=0.1)


def test_letter_paper_of_another_size_sets_the_page_size(owner):
    assert upload(owner, "a5.pdf", letter_paper(size=(419.5, 595.3))).status_code == 201
    reader = PdfReader(io.BytesIO(post(owner, f"/api/orders/{new_order(owner)['id']}/invoice").data))
    assert float(reader.pages[0].mediabox.width) == pytest.approx(419.5, abs=0.1)


def test_the_margins_of_the_paper_can_be_changed_later_and_must_leave_room(owner):
    upload(owner, "p.pdf", letter_paper())
    res = owner.patch("/api/invoice-template", json={"margin_top_mm": 40, "margin_bottom_mm": 20}, headers=H)
    assert res.get_json()["template"]["margin_top_mm"] == 40
    assert owner.patch("/api/invoice-template", json={"margin_top_mm": 201}, headers=H).status_code == 400
    owner.patch("/api/invoice-template", json={"margin_top_mm": 200, "margin_bottom_mm": 200}, headers=H)
    made = post(owner, f"/api/orders/{new_order(owner)['id']}/invoice")
    assert made.status_code == 422 and "marges" in made.get_json()["error"]


# --- the template: upload rules -----------------------------------------------------------------------------------

def encrypted_pdf():
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.encrypt("geheim")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


@pytest.mark.parametrize("name, content", [
    ("t.txt", b"hallo"), ("t.docx", b"geen zip"), ("t.pdf", b"geen pdf"), ("t.docx", b""), ("t.pdf", b"%PDF-1.4 kapot"),
    ("t.docx", b"PK\x03\x04kapot"), ("t.pdf", encrypted_pdf()),
])
def test_a_template_that_cannot_be_used_is_refused(owner, name, content):
    assert upload(owner, name, content).status_code == 400
    assert owner.get("/api/invoice-template").get_json()["template"] is None


def test_a_missing_file_is_refused(owner):
    assert owner.post("/api/invoice-template", data={}, headers=H, content_type="multipart/form-data").status_code == 400


def test_a_template_that_unpacks_to_far_too_much_is_refused(owner):
    import zipfile
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (invoice_docs.MAX_UNPACKED_BYTES + 1))
    assert upload(owner, "bom.docx", out.getvalue()).status_code == 400


def test_a_new_template_replaces_the_old_one_and_its_file(owner, app):
    from pathlib import Path
    upload(owner, "een.pdf", letter_paper("EEN"))
    folder = Path(app.config["APP_CONFIG"].upload_dir)
    first = {p.name for p in folder.glob("invoice-template-*")}
    upload(owner, "twee.pdf", letter_paper("TWEE"))
    second = {p.name for p in folder.glob("invoice-template-*")}
    assert len(first) == len(second) == 1 and first != second
    assert owner.get("/api/invoice-template").get_json()["template"]["filename"] == "twee.pdf"


def test_the_template_can_be_downloaded_and_removed_to_return_to_the_built_in_layout(owner, app):
    from pathlib import Path
    paper = letter_paper("BLIJFT")
    upload(owner, "p.pdf", paper)
    assert owner.get("/api/invoice-template/download").data == paper
    assert owner.delete("/api/invoice-template", headers=H).status_code == 200
    assert owner.get("/api/invoice-template").get_json()["template"] is None
    assert owner.get("/api/invoice-template/download").status_code == 404
    assert list(Path(app.config["APP_CONFIG"].upload_dir).glob("invoice-template-*")) == []
    business(owner)
    assert "Industrieweg 5" in pdf_text(post(owner, f"/api/orders/{new_order(owner)['id']}/invoice").data)


def test_a_template_file_that_went_missing_is_said_so(owner, app):
    from pathlib import Path
    upload(owner, "p.pdf", letter_paper())
    for path in Path(app.config["APP_CONFIG"].upload_dir).glob("invoice-template-*"):
        path.unlink()
    made = post(owner, f"/api/orders/{new_order(owner)['id']}/invoice")
    assert made.status_code == 422 and "opnieuw" in made.get_json()["error"]


# --- money in words -------------------------------------------------------------------------------------------------

@pytest.mark.parametrize("value, lang, shown", [("1234.5", "nl", "€ 1.234,50"), ("1234.5", "en", "€ 1,234.50"), ("0", "nl", "€ 0,00"),
                                                  ("0.005", "nl", "€ 0,01"), ("1000000", "nl", "€ 1.000.000,00")])
def test_amounts_are_written_the_way_the_language_does(value, lang, shown):
    assert invoice_docs.money(value, lang) == shown


# --- the lines add up exactly ---------------------------------------------------------------------------------------

def test_the_net_and_vat_of_the_lines_add_up_to_the_totals_to_the_cent():
    import random
    from app.orders import compute_totals
    rng = random.Random(7)
    for _ in range(300):
        lines = [{"quantity": rng.randint(1, 5), "unit_price": Decimal(rng.randint(1, 9999)) / 100, "vat_pct": rng.choice(["0", "9", "21", "21", "12.5"])}
                 for _ in range(rng.randint(1, 7))]
        totals = compute_totals(lines)
        nets, vats, gross = (list(map(Decimal, totals[k])) for k in ("line_net", "line_vat", "line_totals"))
        assert sum(nets) == Decimal(totals["total_excl"]) and sum(vats) == Decimal(totals["total_vat"]) and sum(gross) == Decimal(totals["total"]), lines
        assert all(n + v == g and 0 <= n <= g for n, v, g in zip(nets, vats, gross)), lines


def test_three_lines_of_99_cents_share_the_cents_of_the_group(owner):
    order = new_order(owner, [line("A", 1, "0.99", "21"), line("B", 1, "0.99", "21"), line("C", 1, "0.99", "21")])
    totals = order["totals"]
    assert sorted(totals["line_net"]) == ["0.81", "0.82", "0.82"] and totals["total_excl"] == "2.45"      # not 3 x 0.82
    assert sum(map(Decimal, totals["line_vat"])) == Decimal("0.52")


# --- the new order and company fields ------------------------------------------------------------------------------------

def test_an_order_keeps_its_customer_number_and_subject(owner):
    order = new_order(owner)
    assert (order["customer_number"], order["subject"]) == ("K9999", "Rack-houders voor oktober")
    bare = owner.post("/api/orders", json={"customer_name": "Piet"}, headers=H).get_json()
    assert (bare["customer_number"], bare["subject"]) == ("", "")
    for bad in ({"customer_number": "x" * 41}, {"subject": "x" * 201}):
        assert owner.post("/api/orders", json={"customer_name": "Piet", **bad}, headers=H).status_code == 400


def test_the_company_details_have_a_website_a_bic_and_a_qr_switch(owner):
    saved = business(owner)
    assert saved["website"] == "flux76.nl" and saved["bic"] == "ABNANL2A" and saved["invoice_qr"] is True
    assert business(owner, invoice_qr=False)["invoice_qr"] is False
    for bad in ("yes", 1, None):
        assert owner.put("/api/preferences", json={"business": {"name": "x", "invoice_qr": bad}}, headers=H).status_code == 400, bad
    assert owner.put("/api/preferences", json={"business": {"name": "x"}}, headers=H).get_json()["business"]["invoice_qr"] is True      # default: on


# --- the payment QR code ----------------------------------------------------------------------------------------------

def test_the_payment_qr_is_a_sepa_code_with_the_amount_and_the_invoice_number():
    order = {"invoice_number": "2026-0001"}
    payload = invoice_docs.epc_payload({"name": "Flux76 Prints", "iban": "nl91 abna 0417 1643 00", "bic": "abnanl2a"}, order, Decimal("29.2"))
    assert payload.split("\n") == ["BCD", "002", "1", "SCT", "ABNANL2A", "Flux76 Prints", "NL91ABNA0417164300", "EUR29.20", "", "", "2026-0001", ""]
    assert invoice_docs.qr_png(payload).startswith(b"\x89PNG")


@pytest.mark.parametrize("business_, amount", [
    ({"name": "Flux", "iban": ""}, "10"), ({"name": "Flux", "iban": "geen iban"}, "10"), ({"name": "", "iban": "NL91ABNA0417164300"}, "10"),
    ({"name": "Flux", "iban": "NL91ABNA0417164300", "invoice_qr": False}, "10"), ({"name": "Flux", "iban": "NL91ABNA0417164300"}, "0"),
])
def test_there_is_no_qr_without_what_a_payment_needs(business_, amount):
    assert invoice_docs.epc_payload(business_, {"invoice_number": "2026-0001"}, Decimal(amount)) is None


def pdf_images(content: bytes) -> int:
    return len(PdfReader(io.BytesIO(content)).pages[0].images)


def test_the_qr_is_on_the_invoice_only_when_it_can_and_should_be(owner):
    business(owner)
    order = new_order(owner)
    assert pdf_images(post(owner, f"/api/orders/{order['id']}/invoice").data) == 1                       # the QR
    business(owner, invoice_qr=False)
    assert pdf_images(post(owner, f"/api/orders/{order['id']}/invoice").data) == 0
    business(owner, iban="")
    assert pdf_images(post(owner, f"/api/orders/{order['id']}/invoice").data) == 0


# --- the logo -----------------------------------------------------------------------------------------------------------

def picture(size=(600, 200), fmt="PNG", colour=(200, 30, 30)):
    from PIL import Image
    out = io.BytesIO()
    Image.new("RGB", size, colour).save(out, fmt)
    return out.getvalue()


def upload_logo(owner, name, content):
    return owner.post("/api/invoice-logo", data={"file": (io.BytesIO(content), name)}, headers=H, content_type="multipart/form-data")


def test_a_logo_can_be_uploaded_looked_at_replaced_and_removed(owner, app):
    from pathlib import Path
    assert owner.get("/api/invoice-logo").get_json() == {"logo": None} and owner.get("/api/invoice-logo/file").status_code == 404
    png = picture()
    res = upload_logo(owner, "mijn logo.png", png)
    assert res.status_code == 201 and res.get_json()["logo"] | {"uploaded_at": 0} == {
        "filename": "mijn logo.png", "kind": "png", "size_bytes": len(png), "width_px": 600, "height_px": 200, "uploaded_at": 0}
    shown = owner.get("/api/invoice-logo/file")
    assert shown.data == png and shown.mimetype == "image/png"
    jpeg = picture((300, 100), "JPEG")
    assert upload_logo(owner, "logo.jpg", jpeg).get_json()["logo"]["kind"] == "jpeg"
    assert owner.get("/api/invoice-logo/file").mimetype == "image/jpeg"
    folder = Path(app.config["APP_CONFIG"].upload_dir)
    assert len(list(folder.glob("invoice-logo-*"))) == 1                                                 # the replaced file is gone
    assert owner.delete("/api/invoice-logo", headers=H).status_code == 200
    assert owner.get("/api/invoice-logo").get_json() == {"logo": None} and list(folder.glob("invoice-logo-*")) == []


@pytest.mark.parametrize("name, content", [
    ("logo.png", b"dit is geen plaatje"), ("logo.svg", b"<svg xmlns='http://www.w3.org/2000/svg'><rect width='9' height='9'/></svg>"),
    ("logo.png", picture(fmt="GIF")), ("logo.png", picture(fmt="BMP")), ("logo.png", picture()[:60]), ("logo.png", b""),
    ("logo.png", b"\x89PNG\r\n\x1a\n" + b"0" * (2 * 1024 * 1024 + 1)),
])
def test_a_file_that_is_not_a_png_or_jpeg_logo_is_refused(owner, name, content):
    res = upload_logo(owner, name, content)
    assert res.status_code == 400 and owner.get("/api/invoice-logo").get_json() == {"logo": None}


def test_an_image_with_far_too_many_pixels_is_refused(owner):
    assert upload_logo(owner, "groot.png", picture((6000, 5000))).status_code == 400


def test_the_logo_needs_a_session(client):
    for method, path in (("get", "/api/invoice-logo"), ("post", "/api/invoice-logo"), ("delete", "/api/invoice-logo"), ("get", "/api/invoice-logo/file")):
        assert getattr(client, method)(path, headers=H).status_code == 401, path


def test_the_logo_is_drawn_on_the_built_in_invoice_but_not_on_letter_paper(owner):
    business(owner)
    upload_logo(owner, "logo.png", picture())
    order = new_order(owner)
    assert pdf_images(post(owner, f"/api/orders/{order['id']}/invoice").data) == 2                      # the logo and the QR
    upload(owner, "paper.pdf", letter_paper())                                                          # the paper has its own header
    assert pdf_images(post(owner, f"/api/orders/{order['id']}/invoice").data) == 1


def test_a_logo_that_went_missing_does_not_stop_the_invoice(owner, app):
    from pathlib import Path
    upload_logo(owner, "logo.png", picture())
    for path in Path(app.config["APP_CONFIG"].upload_dir).glob("invoice-logo-*"):
        path.unlink()
    assert post(owner, f"/api/orders/{new_order(owner)['id']}/invoice").status_code == 200


def test_the_logo_and_the_qr_become_pictures_in_a_word_template(owner):
    business(owner)
    upload(owner, "sjabloon.docx", owner.get("/api/invoice-template/sample").data)
    order = new_order(owner)
    assert len(Document(io.BytesIO(post(owner, f"/api/orders/{order['id']}/invoice").data)).inline_shapes) == 1
    upload_logo(owner, "logo.png", picture((900, 300)))
    made = Document(io.BytesIO(post(owner, f"/api/orders/{order['id']}/invoice").data))
    shapes = made.inline_shapes
    assert len(shapes) == 2
    widths = sorted(round(s.width.cm, 1) for s in shapes)
    assert widths == [2.4, 5.0] and round(max(shapes, key=lambda s: s.width).height.cm, 1) == 1.7      # the QR is square, the logo keeps its shape


def test_a_picture_placeholder_without_a_picture_just_disappears(owner):
    business(owner, iban="")
    upload(owner, "sjabloon.docx", owner.get("/api/invoice-template/sample").data)
    made = post(owner, f"/api/orders/{new_order(owner)['id']}/invoice").data
    assert "{{" not in docx_text(made) and len(Document(io.BytesIO(made)).inline_shapes) == 0
